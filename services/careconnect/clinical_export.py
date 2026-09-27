"""Formatted PDF and Word exports for CareConnect clinical data."""

from __future__ import annotations

from datetime import datetime, timezone
from html import escape
from io import BytesIO
import json
from typing import Any, Iterable


CARECONNECT_BLUE = "2563EB"
HEADING_BLUE = "2E74B5"
DARK_BLUE = "0B2545"
MUTED = "64748B"
LIGHT_FILL = "F2F4F7"
CAUTION_FILL = "FFF7E6"


def _text(value: Any, fallback: str = "Not recorded") -> str:
    if value is None or value == "":
        return fallback
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _date(value: Any) -> str:
    if not value:
        return "Not recorded"
    raw = str(value)
    try:
        parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
        return parsed.strftime("%b %d, %Y %I:%M %p")
    except (TypeError, ValueError):
        return raw.replace("T", " ")


def _medicine_names(prescription: dict) -> str:
    names = [
        _text(item.get("medicine_name"), "")
        for item in prescription.get("medicines", [])
        if item.get("medicine_name")
    ]
    return ", ".join(names) or "Prescription"


def _items(values: Any) -> list[str]:
    if not values:
        return []
    if not isinstance(values, (list, tuple)):
        values = [values]
    return [_text(value, "") for value in values if _text(value, "").strip()]


def _summary_generated_at(payload: dict) -> str:
    return _date(payload.get("exported_at") or datetime.now(timezone.utc).isoformat())


def _set_docx_font(run, *, name="Calibri", size=None, color=None, bold=None, italic=None):
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _configure_docx(document, title: str):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("1F2937")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, HEADING_BLUE, 16, 8),
        "Heading 2": (13, HEADING_BLUE, 12, 6),
        "Heading 3": (12, "1F4D78", 8, 4),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.10

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header.add_run("CareConnect Pro  |  Clinical Export")
    _set_docx_font(header_run, size=9, color=MUTED, bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Confidential clinical information  |  Page ")
    _set_docx_font(footer_run, size=8, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    field_run = OxmlElement("w:r")
    field_text = OxmlElement("w:t")
    field_text.text = "1"
    field_run.append(field_text)
    field.append(field_run)
    footer._p.append(field)

    document.core_properties.title = title
    document.core_properties.author = "CareConnect Pro"
    document.core_properties.subject = "Clinical information export"
    document.core_properties.keywords = "CareConnect, clinical summary, medical records"


def _docx_title(document, title: str, subtitle: str):
    from docx.shared import Pt

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(title)
    _set_docx_font(run, size=24, color=DARK_BLUE, bold=True)

    subtitle_paragraph = document.add_paragraph()
    subtitle_paragraph.paragraph_format.space_after = Pt(16)
    subtitle_run = subtitle_paragraph.add_run(subtitle)
    _set_docx_font(subtitle_run, size=11, color=MUTED)


def _shade_docx_cell(cell, fill: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_docx_cell_width(cell, width_dxa: int):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    width = properties.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        properties.append(width)
    width.set(qn("w:w"), str(width_dxa))
    width.set(qn("w:type"), "dxa")


def _configure_docx_table(table, widths_dxa: list[int], *, header=True):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    table.autofit = False
    properties = table._tbl.tblPr
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")

    width = properties.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        properties.append(width)
    width.set(qn("w:w"), str(sum(widths_dxa)))
    width.set(qn("w:type"), "dxa")

    indent = properties.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")

    cell_margins = properties.find(qn("w:tblCellMar"))
    if cell_margins is None:
        cell_margins = OxmlElement("w:tblCellMar")
        properties.append(cell_margins)
    for side, value in (
        ("top", 80),
        ("start", 120),
        ("bottom", 80),
        ("end", 120),
    ):
        margin = cell_margins.find(qn(f"w:{side}"))
        if margin is None:
            margin = OxmlElement(f"w:{side}")
            cell_margins.append(margin)
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for column_width in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(column_width))
        grid.append(column)

    for row_index, row in enumerate(table.rows):
        if header and row_index == 0:
            row_properties = row._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            row_properties.append(repeat)
        for column_index, cell in enumerate(row.cells):
            _set_docx_cell_width(cell, widths_dxa[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if header and row_index == 0:
                _shade_docx_cell(cell, LIGHT_FILL)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    _set_docx_font(
                        run,
                        size=9.5,
                        color=DARK_BLUE if header and row_index == 0 else "1F2937",
                        bold=header and row_index == 0,
                    )


def _docx_label_value_table(document, rows: Iterable[tuple[str, str]]):
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        _shade_docx_cell(cells[0], LIGHT_FILL)
        for run in cells[0].paragraphs[0].runs:
            run.bold = True
    _configure_docx_table(table, [2700, 6660], header=False)
    for row in table.rows:
        for run in row.cells[0].paragraphs[0].runs:
            _set_docx_font(run, size=9.5, color=DARK_BLUE, bold=True)
    return table


def _docx_bullets(document, values: Any, empty: str = "No items recorded."):
    items = _items(values)
    if not items:
        document.add_paragraph(empty)
        return
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def build_clinical_summary_docx(payload: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    patient = payload.get("patient") or {}
    patient_name = _text(patient.get("name"), "Patient")
    _configure_docx(document, "CareConnect Clinical Summary")
    _docx_title(
        document,
        "Clinical Summary",
        f"{patient_name}  |  Generated {_summary_generated_at(payload)}",
    )

    document.add_heading("Patient overview", level=1)
    blood_pressure = (
        f"{_text(patient.get('systolic_bp'))}/{_text(patient.get('diastolic_bp'))} mmHg"
        if patient.get("systolic_bp") or patient.get("diastolic_bp")
        else "Not recorded"
    )
    _docx_label_value_table(
        document,
        [
            ("Full name", patient_name),
            ("Email", _text(patient.get("email"))),
            ("Age / Gender", f"{_text(patient.get('age'))} / {_text(patient.get('gender'))}"),
            ("Blood type", _text(patient.get("blood_type"))),
            ("Health status", _text(patient.get("status"))),
            ("BMI", _text(patient.get("bmi"))),
            ("Blood pressure", blood_pressure),
            ("Exported by", _text(payload.get("exported_by"))),
        ],
    )

    records = payload.get("records") or []
    document.add_heading(f"Medical records ({len(records)})", level=1)
    if not records:
        document.add_paragraph("No medical records are included in this summary.")
    for record in records:
        document.add_heading(_text(record.get("name"), "Medical record"), level=2)
        meta = document.add_paragraph()
        meta_run = meta.add_run(
            f"{_text(record.get('category'))}  |  "
            f"{_date(record.get('source_date') or record.get('uploaded_at'))}"
        )
        _set_docx_font(meta_run, size=9.5, color=MUTED, bold=True)
        if record.get("analysis_summary"):
            document.add_paragraph(_text(record.get("analysis_summary")))
        if record.get("key_findings"):
            heading = document.add_paragraph()
            heading.paragraph_format.space_after = Pt(2)
            run = heading.add_run("Key findings")
            _set_docx_font(run, size=10.5, color=DARK_BLUE, bold=True)
            _docx_bullets(document, record.get("key_findings"))

    prescriptions = payload.get("prescriptions") or []
    document.add_heading(f"Prescriptions ({len(prescriptions)})", level=1)
    if not prescriptions:
        document.add_paragraph("No prescriptions are included in this summary.")
    for prescription in prescriptions:
        document.add_heading(_medicine_names(prescription), level=2)
        _docx_label_value_table(
            document,
            [
                ("Diagnosis", _text(prescription.get("diagnosis"))),
                ("Status", _text(prescription.get("status"))),
                ("Prescribed by", _text(prescription.get("clinician_name") or prescription.get("clinician_email"))),
                ("Created", _date(prescription.get("created_at"))),
            ],
        )
        medicines = prescription.get("medicines") or []
        if medicines:
            table = document.add_table(rows=1, cols=5)
            table.style = "Table Grid"
            headers = ["Medicine", "Dose", "Frequency", "Duration", "Instructions"]
            for index, label in enumerate(headers):
                table.rows[0].cells[index].text = label
            for medicine in medicines:
                cells = table.add_row().cells
                values = [
                    medicine.get("medicine_name"),
                    medicine.get("dosage"),
                    medicine.get("frequency"),
                    medicine.get("duration"),
                    medicine.get("instructions"),
                ]
                for index, value in enumerate(values):
                    cells[index].text = _text(value, "-")
            _configure_docx_table(table, [1900, 1200, 1500, 1200, 3560])

    appointments = payload.get("appointments") or []
    document.add_heading(f"Appointments ({len(appointments)})", level=1)
    if appointments:
        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        headers = ["Date", "Time", "Type", "Status", "Reason"]
        for index, label in enumerate(headers):
            table.rows[0].cells[index].text = label
        for appointment in appointments:
            cells = table.add_row().cells
            values = [
                appointment.get("date"),
                appointment.get("time"),
                str(appointment.get("type") or "").replace("_", " ").title(),
                str(appointment.get("status") or "").title(),
                appointment.get("reason"),
            ]
            for index, value in enumerate(values):
                cells[index].text = _text(value, "-")
        _configure_docx_table(table, [1300, 950, 1300, 1200, 4610])
    else:
        document.add_paragraph("No appointments are included in this summary.")

    measurements = payload.get("measurement_history") or []
    document.add_heading(f"Body composition and vitals ({len(measurements)})", level=1)
    if measurements:
        table = document.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        headers = ["Recorded", "Weight", "BMI", "Body fat", "Blood pressure", "Status"]
        for index, label in enumerate(headers):
            table.rows[0].cells[index].text = label
        for item in measurements:
            cells = table.add_row().cells
            blood_pressure = (
                f"{_text(item.get('systolic_bp'), '-')}/{_text(item.get('diastolic_bp'), '-')}"
            )
            values = [
                _date(item.get("recorded_at")),
                f"{_text(item.get('weight_kg'), '-')} kg",
                item.get("bmi"),
                f"{_text(item.get('body_fat_percentage'), '-')}%",
                blood_pressure,
                item.get("status"),
            ]
            for index, value in enumerate(values):
                cells[index].text = _text(value, "-")
        _configure_docx_table(table, [1900, 1300, 900, 1200, 1800, 2260])
    else:
        document.add_paragraph("No longitudinal measurements are included in this summary.")

    disclaimer = document.add_paragraph()
    disclaimer.paragraph_format.space_before = Pt(16)
    disclaimer.paragraph_format.space_after = Pt(0)
    disclaimer_run = disclaimer.add_run(
        "Important: " + _text(payload.get("disclaimer"))
    )
    _set_docx_font(disclaimer_run, size=9, color="7A5A00", italic=True)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf_styles():
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch

    styles = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "ClinicalTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=24,
            leading=28,
            textColor=colors.HexColor(f"#{DARK_BLUE}"),
            alignment=TA_LEFT,
            spaceAfter=4,
        ),
        "subtitle": ParagraphStyle(
            "ClinicalSubtitle",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor(f"#{MUTED}"),
            spaceAfter=16,
        ),
        "h1": ParagraphStyle(
            "ClinicalH1",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=16,
            leading=19,
            textColor=colors.HexColor(f"#{HEADING_BLUE}"),
            spaceBefore=16,
            spaceAfter=8,
            keepWithNext=True,
        ),
        "h2": ParagraphStyle(
            "ClinicalH2",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=15,
            textColor=colors.HexColor("#1F4D78"),
            spaceBefore=10,
            spaceAfter=4,
            keepWithNext=True,
        ),
        "body": ParagraphStyle(
            "ClinicalBody",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor("#1F2937"),
            spaceAfter=6,
        ),
        "meta": ParagraphStyle(
            "ClinicalMeta",
            parent=styles["BodyText"],
            fontName="Helvetica-Bold",
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor(f"#{MUTED}"),
            spaceAfter=5,
        ),
        "bullet": ParagraphStyle(
            "ClinicalBullet",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor("#1F2937"),
            leftIndent=0.25 * inch,
            firstLineIndent=-0.12 * inch,
            bulletIndent=0.05 * inch,
            spaceAfter=3,
        ),
        "small": ParagraphStyle(
            "ClinicalSmall",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=8,
            leading=10,
            textColor=colors.HexColor(f"#{MUTED}"),
        ),
        "disclaimer": ParagraphStyle(
            "ClinicalDisclaimer",
            parent=styles["BodyText"],
            fontName="Helvetica-Oblique",
            fontSize=8.5,
            leading=12,
            textColor=colors.HexColor("#7A5A00"),
            borderColor=colors.HexColor("#F4D58D"),
            borderWidth=0.5,
            borderPadding=8,
            backColor=colors.HexColor(f"#{CAUTION_FILL}"),
            spaceBefore=14,
        ),
    }


def _pdf_paragraph(value: Any, style):
    from reportlab.platypus import Paragraph

    return Paragraph(escape(_text(value)).replace("\n", "<br/>"), style)


def _pdf_table(rows: list[list[Any]], widths, styles, *, repeat_rows=1):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle

    converted = [
        [_pdf_paragraph(cell, styles["small"]) for cell in row]
        for row in rows
    ]
    table = Table(converted, colWidths=widths, repeatRows=repeat_rows, hAlign="LEFT")
    commands = [
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{LIGHT_FILL}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor(f"#{DARK_BLUE}")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]
    table.setStyle(TableStyle(commands))
    return table


def _pdf_label_value_table(rows: Iterable[tuple[str, str]], styles):
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    from reportlab.platypus import Table, TableStyle

    converted = [
        [
            _pdf_paragraph(label, styles["small"]),
            _pdf_paragraph(value, styles["small"]),
        ]
        for label, value in rows
    ]
    table = Table(converted, colWidths=[1.75 * inch, 4.75 * inch], hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor(f"#{LIGHT_FILL}")),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return table


def _pdf_page(canvas, document):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter

    canvas.saveState()
    width, height = letter
    canvas.setStrokeColor(colors.HexColor("#CBD5E1"))
    canvas.setLineWidth(0.5)
    canvas.line(document.leftMargin, height - 42, width - document.rightMargin, height - 42)
    canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
    canvas.setFont("Helvetica-Bold", 8)
    canvas.drawString(document.leftMargin, height - 34, "CareConnect Pro  |  Clinical Export")
    canvas.setFont("Helvetica", 8)
    canvas.drawRightString(width - document.rightMargin, 28, f"Confidential  |  Page {document.page}")
    canvas.restoreState()


def build_clinical_summary_pdf(payload: dict) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=0.78 * inch,
        bottomMargin=0.65 * inch,
        title="CareConnect Clinical Summary",
        author="CareConnect Pro",
    )
    styles = _pdf_styles()
    story = []
    patient = payload.get("patient") or {}
    patient_name = _text(patient.get("name"), "Patient")
    story.append(Paragraph("Clinical Summary", styles["title"]))
    story.append(
        Paragraph(
            escape(f"{patient_name}  |  Generated {_summary_generated_at(payload)}"),
            styles["subtitle"],
        )
    )

    story.append(Paragraph("Patient overview", styles["h1"]))
    blood_pressure = (
        f"{_text(patient.get('systolic_bp'))}/{_text(patient.get('diastolic_bp'))} mmHg"
        if patient.get("systolic_bp") or patient.get("diastolic_bp")
        else "Not recorded"
    )
    story.append(
        _pdf_label_value_table(
            [
                ("Full name", patient_name),
                ("Email", _text(patient.get("email"))),
                ("Age / Gender", f"{_text(patient.get('age'))} / {_text(patient.get('gender'))}"),
                ("Blood type", _text(patient.get("blood_type"))),
                ("Health status", _text(patient.get("status"))),
                ("BMI", _text(patient.get("bmi"))),
                ("Blood pressure", blood_pressure),
                ("Exported by", _text(payload.get("exported_by"))),
            ],
            styles,
        )
    )

    records = payload.get("records") or []
    story.append(Paragraph(f"Medical records ({len(records)})", styles["h1"]))
    if not records:
        story.append(Paragraph("No medical records are included in this summary.", styles["body"]))
    for record in records:
        story.append(_pdf_paragraph(record.get("name") or "Medical record", styles["h2"]))
        story.append(
            _pdf_paragraph(
                f"{_text(record.get('category'))}  |  {_date(record.get('source_date') or record.get('uploaded_at'))}",
                styles["meta"],
            )
        )
        if record.get("analysis_summary"):
            story.append(_pdf_paragraph(record.get("analysis_summary"), styles["body"]))
        for finding in _items(record.get("key_findings")):
            story.append(Paragraph(f"- {escape(finding)}", styles["bullet"]))

    prescriptions = payload.get("prescriptions") or []
    story.append(Paragraph(f"Prescriptions ({len(prescriptions)})", styles["h1"]))
    if not prescriptions:
        story.append(Paragraph("No prescriptions are included in this summary.", styles["body"]))
    for prescription in prescriptions:
        story.append(_pdf_paragraph(_medicine_names(prescription), styles["h2"]))
        story.append(
            _pdf_label_value_table(
                [
                    ("Diagnosis", _text(prescription.get("diagnosis"))),
                    ("Status", _text(prescription.get("status"))),
                    ("Prescribed by", _text(prescription.get("clinician_name") or prescription.get("clinician_email"))),
                    ("Created", _date(prescription.get("created_at"))),
                ],
                styles,
            )
        )
        medicines = prescription.get("medicines") or []
        if medicines:
            rows = [["Medicine", "Dose", "Frequency", "Duration", "Instructions"]]
            rows.extend(
                [
                    medicine.get("medicine_name"),
                    medicine.get("dosage"),
                    medicine.get("frequency"),
                    medicine.get("duration"),
                    medicine.get("instructions"),
                ]
                for medicine in medicines
            )
            story.append(_pdf_table(rows, [1.3 * inch, 0.8 * inch, 1.0 * inch, 0.8 * inch, 2.6 * inch], styles))

    appointments = payload.get("appointments") or []
    story.append(Paragraph(f"Appointments ({len(appointments)})", styles["h1"]))
    if appointments:
        rows = [["Date", "Time", "Type", "Status", "Reason"]]
        rows.extend(
            [
                item.get("date"),
                item.get("time"),
                str(item.get("type") or "").replace("_", " ").title(),
                str(item.get("status") or "").title(),
                item.get("reason"),
            ]
            for item in appointments
        )
        story.append(_pdf_table(rows, [0.9 * inch, 0.65 * inch, 0.9 * inch, 0.8 * inch, 3.25 * inch], styles))
    else:
        story.append(Paragraph("No appointments are included in this summary.", styles["body"]))

    measurements = payload.get("measurement_history") or []
    story.append(Paragraph(f"Body composition and vitals ({len(measurements)})", styles["h1"]))
    if measurements:
        rows = [["Recorded", "Weight", "BMI", "Body fat", "Blood pressure", "Status"]]
        rows.extend(
            [
                _date(item.get("recorded_at")),
                f"{_text(item.get('weight_kg'), '-')} kg",
                item.get("bmi"),
                f"{_text(item.get('body_fat_percentage'), '-')}%",
                f"{_text(item.get('systolic_bp'), '-')}/{_text(item.get('diastolic_bp'), '-')}",
                item.get("status"),
            ]
            for item in measurements
        )
        story.append(_pdf_table(rows, [1.3 * inch, 0.85 * inch, 0.6 * inch, 0.8 * inch, 1.15 * inch, 1.8 * inch], styles))
    else:
        story.append(Paragraph("No longitudinal measurements are included in this summary.", styles["body"]))

    story.append(_pdf_paragraph("Important: " + _text(payload.get("disclaimer")), styles["disclaimer"]))
    document.build(story, onFirstPage=_pdf_page, onLaterPages=_pdf_page)
    return output.getvalue()


def _comparison_sections(comparison: dict) -> list[tuple[str, list[str]]]:
    return [
        ("Improved items", _items(comparison.get("improved_items"))),
        ("Worsened items", _items(comparison.get("worsened_items"))),
        ("New concerns", _items(comparison.get("new_concerns") or comparison.get("new_findings"))),
        ("Stable items", _items(comparison.get("stable_items") or comparison.get("common_findings"))),
        ("Recommended next steps", _items(comparison.get("recommended_next_steps"))),
    ]


def build_report_comparison_docx(comparison: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    _configure_docx(document, "CareConnect AI Report Comparison")
    first = comparison.get("first_record") or {}
    second = comparison.get("second_record") or {}
    _docx_title(
        document,
        "AI Report Comparison",
        f"{_text(first.get('name'), 'First report')} compared with {_text(second.get('name'), 'Second report')}",
    )
    _docx_label_value_table(
        document,
        [
            ("First report", f"{_text(first.get('name'))} | {_date(first.get('uploaded_at'))}"),
            ("Second report", f"{_text(second.get('name'))} | {_date(second.get('uploaded_at'))}"),
            ("Generated", _date(datetime.now(timezone.utc).isoformat())),
        ],
    )
    document.add_heading("Comparison summary", level=1)
    document.add_paragraph(_text(comparison.get("ai_summary") or comparison.get("summary")))
    if comparison.get("patient_friendly_explanation"):
        document.add_heading("Patient-friendly explanation", level=2)
        document.add_paragraph(_text(comparison.get("patient_friendly_explanation")))

    metrics = comparison.get("metric_comparison") or []
    document.add_heading(f"Metric comparison ({len(metrics)})", level=1)
    if metrics:
        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for index, label in enumerate(["Metric", "First value", "Second value", "Difference", "Status"]):
            table.rows[0].cells[index].text = label
        for metric in metrics:
            cells = table.add_row().cells
            values = [
                metric.get("metric"),
                metric.get("first_value"),
                metric.get("second_value"),
                metric.get("difference"),
                metric.get("status"),
            ]
            for index, value in enumerate(values):
                cells[index].text = _text(value, "-")
        _configure_docx_table(table, [1900, 1800, 1800, 1400, 2460])
    else:
        document.add_paragraph("No comparable metrics were detected.")

    for heading, items in _comparison_sections(comparison):
        document.add_heading(heading, level=1)
        _docx_bullets(document, items, f"No {heading.lower()} were identified.")

    disclaimer = document.add_paragraph()
    disclaimer.paragraph_format.space_before = Pt(16)
    run = disclaimer.add_run(
        "Important: This comparison is informational and does not replace review by a qualified healthcare professional."
    )
    _set_docx_font(run, size=9, color="7A5A00", italic=True)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def build_report_comparison_pdf(comparison: dict) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import inch
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate

    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=0.78 * inch,
        bottomMargin=0.65 * inch,
        title="CareConnect AI Report Comparison",
        author="CareConnect Pro",
    )
    styles = _pdf_styles()
    story = []
    first = comparison.get("first_record") or {}
    second = comparison.get("second_record") or {}
    story.append(Paragraph("AI Report Comparison", styles["title"]))
    story.append(
        _pdf_paragraph(
            f"{_text(first.get('name'), 'First report')} compared with {_text(second.get('name'), 'Second report')}",
            styles["subtitle"],
        )
    )
    story.append(
        _pdf_label_value_table(
            [
                ("First report", f"{_text(first.get('name'))} | {_date(first.get('uploaded_at'))}"),
                ("Second report", f"{_text(second.get('name'))} | {_date(second.get('uploaded_at'))}"),
                ("Generated", _date(datetime.now(timezone.utc).isoformat())),
            ],
            styles,
        )
    )
    story.append(Paragraph("Comparison summary", styles["h1"]))
    story.append(_pdf_paragraph(comparison.get("ai_summary") or comparison.get("summary"), styles["body"]))
    if comparison.get("patient_friendly_explanation"):
        story.append(Paragraph("Patient-friendly explanation", styles["h2"]))
        story.append(_pdf_paragraph(comparison.get("patient_friendly_explanation"), styles["body"]))

    metrics = comparison.get("metric_comparison") or []
    story.append(Paragraph(f"Metric comparison ({len(metrics)})", styles["h1"]))
    if metrics:
        rows = [["Metric", "First value", "Second value", "Difference", "Status"]]
        rows.extend(
            [
                item.get("metric"),
                item.get("first_value"),
                item.get("second_value"),
                item.get("difference"),
                item.get("status"),
            ]
            for item in metrics
        )
        story.append(_pdf_table(rows, [1.3 * inch, 1.15 * inch, 1.15 * inch, 0.9 * inch, 2.0 * inch], styles))
    else:
        story.append(Paragraph("No comparable metrics were detected.", styles["body"]))

    for heading, items in _comparison_sections(comparison):
        section = [Paragraph(heading, styles["h1"])]
        if items:
            for item in items:
                section.append(Paragraph(f"- {escape(item)}", styles["bullet"]))
        else:
            section.append(
                Paragraph(
                    f"No {heading.lower()} were identified.",
                    styles["body"],
                )
            )
        story.append(KeepTogether(section))

    story.append(
        _pdf_paragraph(
            "Important: This comparison is informational and does not replace review by a qualified healthcare professional.",
            styles["disclaimer"],
        )
    )
    document.build(story, onFirstPage=_pdf_page, onLaterPages=_pdf_page)
    return output.getvalue()
